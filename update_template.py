import docx
import os

template_path = r'templates\OT_template.docx'
doc = docx.Document(template_path)

replacements = {
    "ubicado en Dique 2, Puerto de Santa Fe de la Vera Cruz, Provincia de Santa Fe.": "en el domicilio especificado en el pliego de condiciones.",
    "Cláusulas SSMA del contrato (Anexo SSMA RBA)": "Cláusulas SSMA del contrato (Anexos SSMA vigentes)",
    "sistema Raizen Proveedores": "sistema de proveedores del cliente",
    "Todas las facturas, notas de crédito y débito se cargan en el portal: https://proveedores.raizen.com.ar/interaction/": "Todas las facturas, notas de crédito y débito se cargan en el portal de proveedores correspondiente."
}

# Iterate over all paragraphs
for p in doc.paragraphs:
    full_text = p.text
    changed = False
    for old, new in replacements.items():
        if old in full_text:
            full_text = full_text.replace(old, new)
            changed = True
            
    if changed:
        if p.runs:
            p.runs[0].text = full_text
            for run in p.runs[1:]:
                run.text = ""

# Iterate over all tables
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                full_text = p.text
                changed = False
                for old, new in replacements.items():
                    if old in full_text:
                        full_text = full_text.replace(old, new)
                        changed = True
                        
                if changed:
                    if p.runs:
                        p.runs[0].text = full_text
                        for run in p.runs[1:]:
                            run.text = ""

doc.save(template_path)
print("Template updated successfully.")

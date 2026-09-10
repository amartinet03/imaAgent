import docx

template_path = r'templates\OT_template.docx'
doc = docx.Document(template_path)

# Find paragraph 12
insert_idx = -1
for i, p in enumerate(doc.paragraphs):
    if "12. Anexos de la Oferta Técnica" in p.text:
        insert_idx = i
        break

if insert_idx != -1:
    # Insert new paragraph after it. docx doesn't have a direct insert_paragraph_after,
    # but we can insert before the next one.
    next_p = doc.paragraphs[insert_idx + 1]
    new_p = next_p.insert_paragraph_before("{{TEXTO_REQUISITOS_ADICIONALES}}")
    doc.save(template_path)
    print("Added {{TEXTO_REQUISITOS_ADICIONALES}} successfully.")
else:
    print("Section 12 not found.")

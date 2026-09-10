class WordGenerator:
    """
    Generador de Oferta Técnica (OT).
    Toma la estructura base (histórica) y la adapta al pliego actual, considerando las 
    circulares, visitas de terreno y las respuestas recibidas del cliente.
    """
    def __init__(self, llm_client=None):
        self.llm_client = llm_client
        
    def generate_ot_content_with_llm(self, parsed_data: dict) -> dict:
        """
        Usa el LLM para redactar las secciones dinámicas de la Oferta Técnica en el estilo de IMA.
        """
        import json
        import re
        from langchain_core.prompts import PromptTemplate
        from src.core.analyzer import invoke_with_retry
        
        system_prompt = """
        Eres el Gerente de Licitaciones de IMA Servicios Industriales. Tu objetivo es redactar secciones dinámicas de una Oferta Técnica.
        Queda ESTRICTAMENTE PROHIBIDO arrastrar información operativa, técnica, gremial, geográfica o contractual de otras licitaciones pasadas. Todo lo que sea adaptable debe estar diseñado y redactado EXCLUSIVAMENTE en función de los datos extraídos del nuevo pliego.
        Escribe siempre en tercera persona refiriéndote a la empresa como "IMA". Ej: "IMA asume el compromiso de...", "IMA proveerá...".
        
        REGLAS ESTRICTAS DE DISEÑO Y FORMATO (INYECCIÓN EN WORD):
        1. TEXTO LIMPIO: Redacta con un estilo minimalista, profesional y directo. No inventes títulos internos en MAYÚSCULAS ni uses etiquetas raras.
        2. LISTAS Y VIÑETAS: Cada vez que enumeres algo (tareas, perfiles, herramientas), usa viñetas (guiones -). Asegúrate de que quede un salto de línea (\n) entre los ítems para que no quede un muro de texto.
        3. NO USES MARKDOWN: Como este texto se inyectará en texto plano en Word, NO uses tablas Markdown (|---|) ni negritas (**). Para las tablas, simplemente descríbelas como listas.
        4. DATOS DUROS: Extrae exactamente el alcance, horarios y requerimientos. Si no pide algo explícitamente, responde "Conforme a estándares y requerimientos del cliente".
        
        A partir de los siguientes datos extraídos del pliego, debes generar un ÚNICO objeto JSON estricto con los siguientes campos listos para inyectar en Word:
        - "TITULO_OBRA": Extrae o deduce el nombre oficial corto del servicio a cotizar (Ej: "Obras Civiles Menores", "Mantenimiento Mecánico").
        - "TEXTO_LUGAR_PRESTACION": Redacta detalladamente dónde se ejecutará el servicio.
        - "TEXTO_HORARIOS": Redacta los horarios exigidos y guardias.
        - "TEXTO_ALCANCE_GENERAL": Resumen formal y corporativo del alcance general y tareas.
        - "TEXTO_TAREAS_DEFINIDAS": Descripción operativa y lista de tareas (usa viñetas).
        - "TEXTO_MATERIALES_CLIENTE": Redacta detalladamente qué materiales, servicios o insumos proveerá el cliente. Si el pliego menciona baños químicos, obradores, energía o agua, lístalos. Si no dice nada, redacta "Conforme a requerimientos del cliente".
        - "TEXTO_MATERIALES_IMA": Redacta exhaustivamente los consumibles o materiales que proveerá IMA para la obra.
        - "TEXTO_PERFILES_PERSONAL": Define detalladamente las categorías y cantidad de personal exigido para el servicio, ajustado al gremio (usa viñetas).
        - "TEXTO_REQUISITOS_ADICIONALES": Extrae ÚNICAMENTE documentos exigidos para presentar la OFERTA TÉCNICA (fase licitación, previo a ganar el contrato) como CVs, cronogramas, certificaciones. PROHIBIDO incluir documentos post-adjudicación (ej. Pólizas de Seguro, ART, Garantías, Aptitud Médica). Si no pide documentos pre-adjudicación, responde "Conforme a requerimientos del cliente". Genera lista separada por saltos de línea.
        - "TEXTO_ENCUADRE_GREMIAL": Identifica el Convenio Colectivo aplicable (UOCRA, UOM, Petroleros u otro) y ajusta las categorías del personal a dicho convenio.
        - "TEXTO_HERRAMIENTAS_Y_EPP": Genera un listado nuevo y coherente de herramientas y EPP. IMPORTANTE: NO incluyas artículos de soldadura, mantas o carpas si es una obra puramente civil (usa viñetas).
        - "TEXTO_SSMA": Adapta todas las políticas, normativas (ej. sistemas de permisos, reglas de oro, residuos) a las reglas específicas de ESTE cliente.
        - "TEXTO_CONDICIONES_COMERCIALES": Define jornadas, plazos de certificación, modalidades de pago y ajuste de tarifas dictadas en el pliego actual.
        - "TEXTO_COMUNICACION": Define cómo reportará IMA la gestión (reuniones, informes).
        
        Datos extraídos del pliego:
        {tecnicos}
        {costos}
        {legales}
        
        IMPORTANTE: TU RESPUESTA DEBE SER ÚNICA Y EXCLUSIVAMENTE UN OBJETO JSON VÁLIDO. NO ESCRIBAS NADA FUERA DEL JSON.
        """
        
        prompt = PromptTemplate.from_template(system_prompt)
        
        tecnicos_str = json.dumps(parsed_data.get("tecnicos", {}), ensure_ascii=False, indent=2)
        costos_str = json.dumps(parsed_data.get("costos", {}), ensure_ascii=False, indent=2)
        legales_str = json.dumps(parsed_data.get("legales", {}), ensure_ascii=False, indent=2)
        
        prompt_value = prompt.invoke({"tecnicos": tecnicos_str, "costos": costos_str, "legales": legales_str})
        
        try:
            response = invoke_with_retry(prompt_value)
            content = response.content
            
            if isinstance(content, list):
                content = "".join([c.get("text", "") if isinstance(c, dict) else str(c) for c in content])
            elif not isinstance(content, str):
                content = str(content)
                
            # Limpiar posibles bloques de código markdown
            content = content.replace("```json", "").replace("```", "").strip()
            
            match = re.search(r'\{[\s\S]*\}', content)
            if match:
                json_str = match.group(0)
            else:
                json_str = content
                
            try:
                return json.loads(json_str)
            except json.JSONDecodeError as e:
                print(f"Error parseando JSON estricto: {e}. Intentando rescate por Regex...")
                fallback_dict = {}
                keys = ["TITULO_OBRA", "TEXTO_LUGAR_PRESTACION", "TEXTO_HORARIOS", "TEXTO_ALCANCE_GENERAL", 
                        "TEXTO_TAREAS_DEFINIDAS", "TEXTO_MATERIALES_CLIENTE", "TEXTO_MATERIALES_IMA", 
                        "TEXTO_PERFILES_PERSONAL", "TEXTO_REQUISITOS_ADICIONALES",
                        "TEXTO_ENCUADRE_GREMIAL", "TEXTO_HERRAMIENTAS_Y_EPP", 
                        "TEXTO_SSMA", "TEXTO_CONDICIONES_COMERCIALES", "TEXTO_COMUNICACION"]
                
                for k in keys:
                    # Busca el key y captura el texto hasta la próxima comilla doble
                    m = re.search(fr'"{k}"\s*:\s*"([^"]+)"', json_str)
                    if m:
                        # Reemplaza saltos de línea literales
                        fallback_dict[k] = m.group(1).replace("\\n", "\n")
                
                if fallback_dict:
                    return fallback_dict
                return {}
        except Exception as e:
            print(f"Error generando contenido de OT: {e}")
            return {}

    def draft_technical_offer(self, tender_context: dict, template_path: str, output_path: str, metadata: dict = None) -> str:
        """
        Genera el borrador de la oferta técnica inyectando el contenido generado por 
        la IA en una plantilla Word predefinida con variables.
        """
        import docx
        import os
        import datetime
        
        if not os.path.exists(template_path):
            return "Error: Plantilla OT no encontrada."
            
        parsed_data = tender_context.get("parsed_data", {})
        
        # 1. Generar contenido dinámico con LLM
        dynamic_content = self.generate_ot_content_with_llm(parsed_data)
        
        # 2. Configurar variables de reemplazo (LLM + Metadata base)
        metadata = metadata or {}
        cliente_nombre = metadata.get("cliente", "CLIENTE NO ESPECIFICADO")
        planta = metadata.get("planta", "PLANTA NO ESPECIFICADA")
        
        variables_reemplazo = {
            "{{CLIENTE}}": cliente_nombre,
            "{{PLANTA}}": planta,
            "{{PROCESO}}": metadata.get("proceso", "PROCESO NO ESPECIFICADO"),
            "{{NUMERO_LICITACION}}": metadata.get("proceso", "LIC-PENDIENTE"),
            "{{FECHA_EMISION}}": datetime.datetime.now().strftime("%B %Y").capitalize(),
            "{{CLIENTE_CORTO}}": cliente_nombre.split(" ")[0] if cliente_nombre else "CLIENTE",
        }
        
        # Agregar el contenido dinámico del LLM a las variables de reemplazo
        for key in ["TITULO_OBRA", "TEXTO_LUGAR_PRESTACION", "TEXTO_HORARIOS", "TEXTO_ALCANCE_GENERAL", 
                    "TEXTO_TAREAS_DEFINIDAS", "TEXTO_MATERIALES_CLIENTE", "TEXTO_MATERIALES_IMA", 
                    "TEXTO_PERFILES_PERSONAL", "TEXTO_REQUISITOS_ADICIONALES",
                    "TEXTO_ENCUADRE_GREMIAL", "TEXTO_HERRAMIENTAS_Y_EPP", 
                    "TEXTO_SSMA", "TEXTO_CONDICIONES_COMERCIALES", "TEXTO_COMUNICACION"]:
            variables_reemplazo[f"{{{{{key}}}}}"] = dynamic_content.get(key, "A definir según pliego.")

        # 3. Inyectar variables en el documento
        doc = docx.Document(template_path)
        from docx.oxml.ns import qn
        
        # Inyectar Anexos Dinámicos en la Tabla
        req_adic = dynamic_content.get("TEXTO_REQUISITOS_ADICIONALES", "")
        if req_adic and "Conforme" not in req_adic:
            variables_reemplazo["{{TEXTO_REQUISITOS_ADICIONALES}}"] = ""
            if len(doc.tables) > 0:
                annex_table = doc.tables[-1] # La última tabla es la de Anexos
                from copy import deepcopy
                
                if '\n' not in req_adic and ',' in req_adic:
                    items = [it.strip() for it in req_adic.split(',') if len(it.strip()) > 5]
                else:
                    items = [it.replace("-", "").strip() for it in req_adic.split('\n') if it.strip()]
                    
                start_index = len(annex_table.rows) # Incluye encabezado, si hay 8 anexos + 1 header = 9
                for i, item in enumerate(items):
                    if item:
                        # Clonar la última fila para mantener exactamente los bordes y estilos
                        last_row = annex_table.rows[-1]
                        new_row_xml = deepcopy(last_row._tr)
                        annex_table._tbl.append(new_row_xml)
                        
                        new_row = annex_table.rows[-1]
                        new_row.cells[0].text = f"Anexo {start_index + i}"
                        new_row.cells[1].text = item.capitalize()

        # Reemplazo profundo en todo el árbol XML (Párrafos, Tablas, Formas, Cuadros de texto)
        def replace_in_element(element):
            for t in element.iter(qn('w:t')):
                if t.text:
                    for key, val in variables_reemplazo.items():
                        if key in t.text:
                            t.text = t.text.replace(key, str(val))
                            
        replace_in_element(doc.element)
        
        # Reemplazo en Encabezados y Pies de página
        for section in doc.sections:
            for header in [section.header, section.first_page_header, section.even_page_header]:
                if header: replace_in_element(header._element)
            for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
                if footer: replace_in_element(footer._element)

        doc.save(output_path)
        return output_path

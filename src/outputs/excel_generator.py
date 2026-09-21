import datetime
import os
import shutil
import time
import openpyxl


class ExcelGenerator:
    """
    Generador del Excel de Costos dinámico usando la plantilla base limpia.
    Garantiza que ningún valor quede hardcodeado de la licitación de referencia
    y que todas las fórmulas de cálculo se actualicen automáticamente.
    """
    def __init__(self, llm_client=None):
        self.llm_client = llm_client

    def _search_price(self, description: str):
        """
        Busca un precio de referencia en el mercado argentino usando DuckDuckGo
        y lo normaliza con el LLM.
        """
        try:
            from langchain_community.tools import DuckDuckGoSearchRun
            search = DuckDuckGoSearchRun()
            query = f"precio aproximado unitario de {description} en argentina mercado libre o proveedor"
            res = search.invoke(query)

            if self.llm_client:
                prompt = (
                    f"Basado en este resultado de búsqueda web: '{res}', dame SOLO el precio estimado "
                    f"en ARS (Pesos Argentinos) de '{description}'. Devuelve ÚNICAMENTE un número flotante "
                    f"sin símbolos de moneda ni separador de miles (ej: 15400.50). "
                    f"En la segunda línea, pon una referencia muy breve de la fuente."
                )
                llm_res = self.llm_client.invoke(prompt)
                content = llm_res.content if hasattr(llm_res, 'content') else str(llm_res)
                parts = content.strip().split('\n', 1)
                try:
                    price = float(parts[0].replace('$', '').replace(',', '').strip())
                except Exception:
                    price = 0.0
                source = parts[1].strip() if len(parts) > 1 else "Búsqueda web"
                return price, source
            return 0.0, "Búsqueda web (sin IA)"
        except Exception as e:
            return 0.0, f"No disponible ({e})"

    def _unmerge_if_needed(self, ws, coord: str):
        """Descombina celdas si forman parte de un rango combinado para evitar MergedCell errors."""
        for merged_range in list(ws.merged_cells.ranges):
            if coord in merged_range:
                ws.unmerge_cells(str(merged_range))

    def generate_cost_excel(self, tender_context: dict, template_path: str, output_path: str) -> str:
        """
        Genera el libro de costos completo a partir de la plantilla limpia, poblando
        de forma dinámica y segura cada hoja con los datos de la licitación actual.
        """
        parsed_data = tender_context.get("parsed_data", {})
        costos = parsed_data.get("costos", {})
        metadata = parsed_data.get("metadata", {})
        tecnicos = parsed_data.get("tecnicos", {})
        tender_id = tender_context.get("tender_id", "0")

        # Asegurar directorio de destino y copiar plantilla limpia
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        shutil.copy2(template_path, output_path)

        wb = openpyxl.load_workbook(output_path)

        # ---------------------------------------------------------------------
        # 1. Variables Globales de la Licitación
        # ---------------------------------------------------------------------
        cliente = metadata.get("cliente") or "CLIENTE DESCONOCIDO"
        nombre_pliego = metadata.get("nombre_pliego") or metadata.get("proceso") or "LICITACIÓN S/N"
        proceso = metadata.get("proceso") or "S/N"
        alcance = tecnicos.get("alcance_general") or f"Servicio para {cliente} - {proceso}"

        # Duración del contrato en años para la hoja de Cierre y Síntesis Económica
        meses_contrato = costos.get("plazo_contrato_meses")
        if meses_contrato:
            try:
                duracion_anios = max(1, int(round(float(meses_contrato) / 12.0)))
            except Exception:
                duracion_anios = 1
        else:
            duracion_anios = 1

        # ---------------------------------------------------------------------
        # 2. Hoja de CIERRE (Portada y Parámetros Maestros)
        # ---------------------------------------------------------------------
        if "Hoja de CIERRE" in wb.sheetnames:
            ws_cierre = wb["Hoja de CIERRE"]
            ws_cierre["I3"] = "IMA Servicios Industriales Argentina S.A."
            ws_cierre["I5"] = nombre_pliego.upper()
            # La celda I7 alimenta CLIENTE a todas las otras 11 hojas mediante fórmulas
            ws_cierre["I7"] = cliente.upper()
            ws_cierre["AD7"] = datetime.date.today()
            ws_cierre["I9"] = f"Nº LICITACIÓN: {proceso}"
            ws_cierre["I13"] = "IMA Cotizaciones"
            ws_cierre["X13"] = "IMA Ventas"
            ws_cierre["I14"] = "Gerencia Técnica"
            ws_cierre["I16"] = f"PRE-ARG-{tender_id}"
            ws_cierre["B20"] = alcance
            ws_cierre["AD45"] = duracion_anios

        # ---------------------------------------------------------------------
        # 3. 1 MANO DE OBRA (Directa e Indirecta)
        # ---------------------------------------------------------------------
        if "1 MANO DE OBRA" in wb.sheetnames:
            ws_mo = wb["1 MANO DE OBRA"]
            # Resetear todas las horas de mano de obra a 0
            for r in range(14, 24):  # Directa
                ws_mo[f"C{r}"] = 0
            for r in range(26, 36):  # Indirecta
                ws_mo[f"C{r}"] = 0
            for r in range(38, 42):  # Eventual
                ws_mo[f"C{r}"] = 0

            mo_list = costos.get("mano_de_obra", [])
            for item in mo_list:
                if isinstance(item, dict):
                    rol_str = str(item.get("rol", item.get("item", ""))).lower()
                    try:
                        cant = float(item.get("cantidad", 1))
                    except Exception:
                        cant = 1.0
                    try:
                        horas_input = float(item.get("horas_mensuales", 0))
                    except Exception:
                        horas_input = 0.0

                    if horas_input > 0:
                        horas = (cant * horas_input) if horas_input <= 250 else horas_input
                    else:
                        horas = cant * 187.0
                else:
                    rol_str = str(item).lower()
                    cant = 1.0
                    horas = 187.0

                # Clasificación y mapeo a filas
                if "especializ" in rol_str or "soldador" in rol_str or "cañista" in rol_str or "instrumentista" in rol_str:
                    ws_mo["C19"] = float(ws_mo["C19"].value or 0) + horas
                elif "medio oficial" in rol_str:
                    ws_mo["C17"] = float(ws_mo["C17"].value or 0) + horas
                elif "ayudante" in rol_str:
                    ws_mo["C16"] = float(ws_mo["C16"].value or 0) + horas
                elif "oficial" in rol_str or "montador" in rol_str or "mecanico" in rol_str or "electricista" in rol_str:
                    ws_mo["C18"] = float(ws_mo["C18"].value or 0) + horas
                elif "administrador" in rol_str or "gerente" in rol_str or "coordinador" in rol_str:
                    ws_mo["C26"] = float(ws_mo["C26"].value or 0) + horas
                elif "jefe" in rol_str:
                    ws_mo["C27"] = float(ws_mo["C27"].value or 0) + horas
                elif any(k in rol_str for k in ["seguridad", "ssma", "hse", "ehs", "h&se"]):
                    ws_mo["C28"] = float(ws_mo["C28"].value or 0) + horas
                elif "supervisor" in rol_str or "capataz" in rol_str:
                    ws_mo["C29"] = float(ws_mo["C29"].value or 0) + horas
                elif "administrativo" in rol_str:
                    ws_mo["C30"] = float(ws_mo["C30"].value or 0) + horas
                elif "pañol" in rol_str:
                    ws_mo["C31"] = float(ws_mo["C31"].value or 0) + horas
                elif "chofer" in rol_str:
                    ws_mo["C32"] = float(ws_mo["C32"].value or 0) + horas
                elif "limpieza" in rol_str:
                    ws_mo["C33"] = float(ws_mo["C33"].value or 0) + horas
                else:
                    # Fallback general a Oficial
                    ws_mo["C18"] = float(ws_mo["C18"].value or 0) + horas

        # ---------------------------------------------------------------------
        # 4. 2 SUMINISTRO DE MATERIALES
        # Columnas: B=Item, C=Descripción, D=Precio Unit., E=Unidad, F=Cantidad, G=Total (=D*F), H=Fuente
        # ---------------------------------------------------------------------
        if "2 SUMINISTRO DE MATERIALES" in wb.sheetnames:
            ws_mat = wb["2 SUMINISTRO DE MATERIALES"]
            # Limpiar filas 10 a 19 (sección 2.1)
            for r in range(10, 20):
                ws_mat[f"C{r}"] = None
                ws_mat[f"D{r}"] = 0.0
                ws_mat[f"E{r}"] = None
                ws_mat[f"F{r}"] = 0.0
                ws_mat[f"G{r}"] = f"=D{r}*F{r}"
                ws_mat[f"H{r}"] = None

            mat_list = costos.get("suministro_materiales", [])
            r = 10
            for item in mat_list[:10]:
                if isinstance(item, dict):
                    desc = item.get("item", item.get("descripcion", str(item)))
                    try:
                        cant = float(item.get("cantidad", 1))
                    except Exception:
                        cant = 1.0
                    unit = item.get("unidad", "u")
                else:
                    desc = str(item)
                    cant = 1.0
                    unit = "u"

                price, source = self._search_price(desc)
                time.sleep(0.5)

                for col in ['C', 'D', 'E', 'F', 'G', 'H']:
                    self._unmerge_if_needed(ws_mat, f"{col}{r}")

                ws_mat[f"C{r}"] = desc
                ws_mat[f"D{r}"] = price
                ws_mat[f"E{r}"] = unit
                ws_mat[f"F{r}"] = cant
                ws_mat[f"G{r}"] = f"=D{r}*F{r}"
                ws_mat[f"H{r}"] = source
                r += 1

        # ---------------------------------------------------------------------
        # 5. 3 SUMINISTRO DE INSUMOS
        # Columnas: B=Item, C=Descripción, D=Precio Unit., E=Cantidad, F=Total (=+E*D), G=Fuente
        # ---------------------------------------------------------------------
        if "3 SUMINISTRO DE INSUMOS" in wb.sheetnames:
            ws_ins = wb["3 SUMINISTRO DE INSUMOS"]
            # Limpiar filas 10 a 38 (sección 3.1)
            for r in range(10, 39):
                ws_ins[f"C{r}"] = None
                ws_ins[f"D{r}"] = 0.0
                ws_ins[f"E{r}"] = 0.0
                ws_ins[f"F{r}"] = f"=+E{r}*D{r}"
                ws_ins[f"G{r}"] = None

            ins_list = costos.get("suministro_insumos", [])
            r = 10
            for item in ins_list[:28]:
                if isinstance(item, dict):
                    desc = item.get("item", item.get("descripcion", str(item)))
                    try:
                        cant = float(item.get("cantidad", 1))
                    except Exception:
                        cant = 1.0
                else:
                    desc = str(item)
                    cant = 1.0

                price, source = self._search_price(desc)
                time.sleep(0.5)

                for col in ['C', 'D', 'E', 'F', 'G']:
                    self._unmerge_if_needed(ws_ins, f"{col}{r}")

                ws_ins[f"C{r}"] = desc
                ws_ins[f"D{r}"] = price
                ws_ins[f"E{r}"] = cant
                ws_ins[f"F{r}"] = f"=+E{r}*D{r}"
                ws_ins[f"G{r}"] = source
                r += 1

        # ---------------------------------------------------------------------
        # 6. 4 SUBCONTRATACIONES
        # Columnas: B=Item, C=Descripción, D=Precio Unit., E=Cantidad, F=Total (=+E*D), G=Fuente
        # ---------------------------------------------------------------------
        if "4 SUBCONTRATACIONES" in wb.sheetnames:
            ws_sub = wb["4 SUBCONTRATACIONES"]
            # Limpiar filas 10 a 25 (sección 4.1)
            for r in range(10, 26):
                ws_sub[f"C{r}"] = None
                ws_sub[f"D{r}"] = 0.0
                ws_sub[f"E{r}"] = 0.0
                ws_sub[f"F{r}"] = f"=+E{r}*D{r}"
                ws_sub[f"G{r}"] = None

            sub_list = costos.get("subcontrataciones", [])
            r = 10
            for item in sub_list[:16]:
                if isinstance(item, dict):
                    desc = item.get("item", item.get("descripcion", str(item)))
                    try:
                        cant = float(item.get("cantidad", 1))
                    except Exception:
                        cant = 1.0
                else:
                    desc = str(item)
                    cant = 1.0

                price, source = self._search_price(desc)
                time.sleep(0.5)

                for col in ['C', 'D', 'E', 'F', 'G']:
                    self._unmerge_if_needed(ws_sub, f"{col}{r}")

                ws_sub[f"C{r}"] = desc
                ws_sub[f"D{r}"] = price
                ws_sub[f"E{r}"] = cant
                ws_sub[f"F{r}"] = f"=+E{r}*D{r}"
                ws_sub[f"G{r}"] = source
                r += 1

        # ---------------------------------------------------------------------
        # 7. 5 AMORTIZACIÓN (Vehículos y Equipos Móviles)
        # ---------------------------------------------------------------------
        amort_sheet_name = next((s for s in wb.sheetnames if "AMORTIZA" in s.upper()), None)
        if amort_sheet_name:
            ws_am = wb[amort_sheet_name]
            # Limpiar sección 5.1 Rodados (filas 11 a 22)
            for r in range(11, 23):
                ws_am[f"C{r}"] = None
                ws_am[f"D{r}"] = 0.0
                ws_am[f"E{r}"] = 0.0
                ws_am[f"F{r}"] = f"=+D{r}*E{r}"
                ws_am[f"G{r}"] = 5
                ws_am[f"H{r}"] = 0.5
                ws_am[f"J{r}"] = 0.02
                ws_am[f"K{r}"] = f"=PMT(J{r},G{r}*12,F{r}-I{r},0)-I{r}*J{r}"
                ws_am[f"L{r}"] = None

            veh_list = costos.get("amortizacion_vehiculos", [])
            r = 11
            for item in veh_list[:12]:
                if isinstance(item, dict):
                    desc = item.get("item", item.get("vehiculo", str(item)))
                    try:
                        cant = float(item.get("cantidad", 1))
                    except Exception:
                        cant = 1.0
                    try:
                        plazo = float(item.get("plazo_amortizacion", 5))
                    except Exception:
                        plazo = 5.0
                else:
                    desc = str(item)
                    cant = 1.0
                    plazo = 5.0

                price, source = self._search_price(f"vehiculo {desc}")
                time.sleep(0.5)

                for col in ['C', 'D', 'E', 'F', 'G', 'H', 'J', 'K', 'L']:
                    self._unmerge_if_needed(ws_am, f"{col}{r}")

                ws_am[f"C{r}"] = desc
                ws_am[f"D{r}"] = price
                ws_am[f"E{r}"] = cant
                ws_am[f"F{r}"] = f"=+D{r}*E{r}"
                ws_am[f"G{r}"] = plazo
                ws_am[f"H{r}"] = 0.5
                ws_am[f"J{r}"] = 0.02
                ws_am[f"K{r}"] = f"=PMT(J{r},G{r}*12,F{r}-I{r},0)-I{r}*J{r}"
                ws_am[f"L{r}"] = source
                r += 1

        # ---------------------------------------------------------------------
        # 8. Equipos (Herramientas y Equipos Menores)
        # Columnas: A=Item, B=Descripción, C=Cantidad, D=PU, E=Total (=C*D), F=Fuente
        # ---------------------------------------------------------------------
        if "Equipos" in wb.sheetnames:
            ws_eq = wb["Equipos"]
            # Limpiar filas 4 a 125 preservando totalizador en E126
            for r in range(4, 126):
                ws_eq[f"A{r}"] = None
                ws_eq[f"B{r}"] = None
                ws_eq[f"C{r}"] = 0.0
                ws_eq[f"D{r}"] = 0.0
                ws_eq[f"E{r}"] = f"=C{r}*D{r}"
                ws_eq[f"F{r}"] = None
                ws_eq[f"G{r}"] = None

            eq_list = costos.get("equipos_menores", [])
            r = 4
            for idx, item in enumerate(eq_list[:120]):
                if isinstance(item, dict):
                    desc = item.get("item", item.get("equipo", str(item)))
                    try:
                        cant = float(item.get("cantidad", 1))
                    except Exception:
                        cant = 1.0
                else:
                    desc = str(item)
                    cant = 1.0

                price, source = self._search_price(desc)
                time.sleep(0.5)

                for col in ['A', 'B', 'C', 'D', 'E', 'F']:
                    self._unmerge_if_needed(ws_eq, f"{col}{r}")

                ws_eq[f"A{r}"] = idx + 1
                ws_eq[f"B{r}"] = desc
                ws_eq[f"C{r}"] = cant
                ws_eq[f"D{r}"] = price
                ws_eq[f"E{r}"] = f"=C{r}*D{r}"
                ws_eq[f"F{r}"] = source
                r += 1

        # ---------------------------------------------------------------------
        # 9. Guardar libro y finalizar
        # ---------------------------------------------------------------------
        wb.save(output_path)
        return output_path

    def validate_cost_excel(self, tender_context: dict, uploaded_excel_path: str) -> str:
        """
        Lee el Excel completado por el equipo y lo compara contra los requerimientos 
        extraídos del pliego para detectar incongruencias usando el LLM.
        """
        if not self.llm_client:
            raise ValueError("No LLM client configured for Excel validation")

        parsed_data = tender_context.get("parsed_data", {})
        costos = parsed_data.get("costos", {})
        req_text = str(costos)

        wb = openpyxl.load_workbook(uploaded_excel_path, data_only=True)
        excel_dump = ""

        sheets_to_read = [
            "Hoja de CIERRE", "1 MANO DE OBRA", "Equipos", "3 SUMINISTRO DE INSUMOS",
            "Presupuesto Interno", "2 SUMINISTRO DE MATERIALES", "4 SUBCONTRATACIONES"
        ]
        for s in wb.sheetnames:
            if any(name.upper() in s.upper() for name in sheets_to_read) or "AMORTIZA" in s.upper():
                ws = wb[s]
                excel_dump += f"\n--- Pestaña: {s} ---\n"
                for row in ws.iter_rows(min_row=1, max_row=45, values_only=True):
                    if any(c is not None and str(c).strip() not in ['', 'None', '0', '0.0'] for c in row):
                        row_text = " | ".join(str(c) for c in row if c is not None and str(c).strip() != '')
                        excel_dump += row_text + "\n"

        prompt = f"""
Actúa como un controlador de costos y auditor de licitaciones industriales de ingeniería.
A continuación tienes los requisitos y parámetros mínimos extraídos del Pliego:
{req_text}

Y aquí tienes el volcado de datos del Excel de Costos generado/completado:
{excel_dump}

Tu tarea es revisar si hay incongruencias entre lo que exige el pliego y lo que figura cotizado en el Excel:
- Mano de Obra: ¿Se contemplan los roles exigidos (oficiales, supervisores, técnicos de seguridad)?
- Equipos y Vehículos: ¿Están consideradas las camionetas/equipos obligatorios?
- Insumos, EPP y Materiales: ¿Se incluyeron los ítems requeridos?
- Seguros y Subcontratos: ¿Figuran las coberturas mínimas exigidas?

Devuelve un informe ejecutivo en viñetas:
1. Resumen de Consistencia General.
2. Incongruencias o Desvíos Detectados (si los hay).
3. Recomendaciones para el Cotizador.
"""
        response = self.llm_client.invoke(prompt)
        return response.content if hasattr(response, 'content') else str(response)

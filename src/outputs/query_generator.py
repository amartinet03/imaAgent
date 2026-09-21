class QueryGenerator:
    """
    Generador de Consultas (RFI - Request for Information).
    Toma el contexto de la licitación (documentos analizados, incongruencias detectadas)
    y formatea un documento profesional (Word/PDF) con las preguntas definitivas para el cliente.
    """
    def __init__(self, llm_client=None):
        self.llm_client = llm_client
        

    def generate_docx_from_json(self, parsed_data: dict, template_path: str, output_path: str) -> str:
        """
        Genera el documento de consultas formales inyectando los datos JSON en la plantilla Word.
        """
        import docx
        import os
        
        if not os.path.exists(template_path):
            return "Error: Plantilla de Consultas no encontrada."
            
        doc = docx.Document(template_path)
        
        # Reemplazo total de variables en los párrafos usando metadata del LLM
        metadata = parsed_data.get("metadata", {})
        lista_docs = str(metadata.get("lista_documentos", "No especificado en pliego"))
        lista_docs = ", ".join([doc.strip() for doc in lista_docs.split(",") if doc.strip()])
        
        variables_reemplazo = {
            "{{NOMBREPLIEGO}}": str(metadata.get("nombre_pliego", "No especificado en pliego")),
            "{{CLIENTE}}": str(metadata.get("cliente", "No especificado en pliego")),
            "{{PROCESO}}": str(metadata.get("proceso", "No especificado en pliego")),
            "{{MONEDA}}": str(metadata.get("moneda", "No especificado en pliego")),
            "{{PLANTA}}": str(metadata.get("planta", "No especificado en pliego")),
            "{{REQUIRENTE}}": str(metadata.get("requirente", "No especificado en pliego")),
            "{{COMPRADOR}}": str(metadata.get("comprador", "No especificado en pliego")),
            "{{LISTA_DOCUMENTOS}}": lista_docs
        }
        
        for p in doc.paragraphs:
            full_text = p.text
            has_var = any(key in full_text for key in variables_reemplazo.keys())
            has_pampa = "PAMPA" in full_text.upper() and "pampa" not in str(metadata.get("cliente", "")).lower()
            if has_var or has_pampa:
                # Realizar reemplazos en el texto completo
                for key, val in variables_reemplazo.items():
                    full_text = full_text.replace(key, val)
                
                # Resguardo de seguridad: si existe texto remanente de PAMPA, reemplazarlo por el cliente actual
                cliente_val = str(metadata.get("cliente", "el Cliente"))
                for pampa_target in ["PAMPA ENERGÍA S.A.", "PAMPA ENERGIA S.A.", "Pampa Energía S.A."]:
                    if pampa_target in full_text:
                        full_text = full_text.replace(pampa_target, cliente_val)
                
                # Solución a Run Splitting (Opción B):
                # Asignamos el texto reconstruido al primer run para mantener su estilo base,
                # y vaciamos el resto de los runs para evitar duplicados.
                if p.runs:
                    p.runs[0].text = full_text
                    for run in p.runs[1:]:
                        run.text = ""
        
        # Identificar las tablas dinámicamente y vaciar las filas viejas (manteniendo el encabezado)
        tabla_consultas = None
        tabla_inconsistencias = None
        
        for table in doc.tables:
            if not table.rows:
                continue
            headers = [cell.text.lower().strip() for cell in table.rows[0].cells]
            
            if "categoría" in headers or "categoria" in headers:
                tabla_consultas = table
            elif "tipo" in headers or "documentos en conflicto" in headers:
                tabla_inconsistencias = table

        # Función auxiliar para limpiar la tabla dejando solo el encabezado
        def clear_table(table):
            # Iteramos en reversa para no afectar el índice mientras borramos (XML level)
            for row in table.rows[1:][::-1]:
                tbl = table._tbl
                tr = row._tr
                tbl.remove(tr)

        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        def set_cell_content(cell, text: str, is_bold: bool = False, align=None):
            p = cell.paragraphs[0]
            for r in list(p.runs):
                p._p.remove(r._r)
            if align is not None:
                p.alignment = align
            lines = text.split("\n")
            for idx, line in enumerate(lines):
                run = p.add_run(line)
                run.font.name = "Arial"
                run.font.size = Pt(9)
                run.bold = is_bold
                if idx < len(lines) - 1:
                    run_br = p.add_run("\n")
                    run_br.font.name = "Arial"
                    run_br.font.size = Pt(9)

        # Contador global para numeración continua de filas en ambas tablas
        contador_global = 1

        # Llenar Tabla de Consultas Generales
        if tabla_consultas and "consultas_generales" in parsed_data:
            clear_table(tabla_consultas)
            consultas = parsed_data["consultas_generales"]
            
            for c in consultas:
                row_cells = tabla_consultas.add_row().cells
                if len(row_cells) >= 4:
                    set_cell_content(row_cells[0], str(contador_global), is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
                    set_cell_content(row_cells[1], str(c.get("categoria", "")), is_bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)
                    archivo_info = f"{c.get('archivo_origen', '')}\nPág: {c.get('pagina_origen', 'N/A')}\n\"{c.get('cita_textual', 'N/A')}\""
                    set_cell_content(row_cells[2], archivo_info, is_bold=False)
                    set_cell_content(row_cells[3], str(c.get("consulta", "")), is_bold=False)
                    contador_global += 1

        # Llenar Tabla de Inconsistencias
        if tabla_inconsistencias and "inconsistencias" in parsed_data:
            clear_table(tabla_inconsistencias)
            incons = parsed_data["inconsistencias"]
            
            for i in incons:
                row_cells = tabla_inconsistencias.add_row().cells
                if len(row_cells) >= 4:
                    set_cell_content(row_cells[0], str(contador_global), is_bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
                    set_cell_content(row_cells[1], str(i.get("tipo", "Inconsistencia")), is_bold=False, align=WD_ALIGN_PARAGRAPH.CENTER)
                    
                    # Formateo correcto: si es lista, unir con comas
                    docs_conflicto = i.get("documentos_conflicto", "")
                    if isinstance(docs_conflicto, list):
                        docs_conflicto = ", ".join(docs_conflicto)
                    archivo_info = f"{docs_conflicto}\nPág: {i.get('pagina_origen', 'N/A')}\n\"{i.get('cita_textual', 'N/A')}\""
                    set_cell_content(row_cells[2], archivo_info, is_bold=False)
                    set_cell_content(row_cells[3], str(i.get("descripcion_pregunta", "")), is_bold=False)
                    contador_global += 1
                    
        # Al hacer save a 'output_path', mantenemos intacta la plantilla original.
        doc.save(output_path)
        return output_path

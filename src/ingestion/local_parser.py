import os
import pdfplumber
import pandas as pd
from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document

def extract_text_from_file(uploaded_file):
    """
    Extrae texto dependiendo de la extensión del archivo.
    """
    filename = uploaded_file.name.lower() if hasattr(uploaded_file, 'name') else os.path.basename(uploaded_file).lower()
    raw_text = ""
    
    try:
        if filename.endswith(".pdf"):
            with pdfplumber.open(uploaded_file) as pdf:
                for i, page in enumerate(pdf.pages):
                    # Usamos layout=True para mantener en lo posible la estructura tabular original
                    text = page.extract_text(layout=True)
                    if text:
                        # Inyectar marcador de página para trazabilidad
                        raw_text += f"\n--- [Página {i+1}] ---\n"
                        raw_text += text + "\n"
                    
        elif filename.endswith(".docx"):
            doc = DocxDocument(uploaded_file)
            for para in doc.paragraphs:
                raw_text += para.text + "\n"
                
        elif filename.endswith(".xlsx") or filename.endswith(".xls"):
            # Leer todas las hojas
            dfs = pd.read_excel(uploaded_file, sheet_name=None)
            for sheet_name, df in dfs.items():
                raw_text += f"--- Hoja: {sheet_name} ---\n"
                raw_text += df.to_string() + "\n"
        else:
            print(f"Formato no soportado: {filename}")
            
    except Exception as e:
        print(f"Error al procesar {filename}: {e}")
        
    return raw_text

def process_file_to_langchain_docs(uploaded_file, chunk_size=2000, chunk_overlap=400):
    raw_text = extract_text_from_file(uploaded_file)
    
    if not raw_text.strip():
        return []

    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        length_function=len,
        separators=["\n\n", "\n", " ", ""]
    )
    
    chunks = text_splitter.split_text(raw_text)
    
    # Inyectamos el nombre del archivo en la Metadata (Clave para detectar incongruencias por archivo)
    source_name = getattr(uploaded_file, "name", os.path.basename(uploaded_file) if isinstance(uploaded_file, str) else "Desconocido")
    documents = [Document(page_content=chunk, metadata={"source": source_name}) for chunk in chunks]
    
    return documents
